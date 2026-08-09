"""Generate AML-2203 Final Term Project Word (.docx) technical report."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "AML-2203-Code-Review-Technical-Report.docx"
BEFORE = (ROOT / "before" / "update_knowledge_before.py").read_text(encoding="utf-8")
AFTER = (ROOT / "after" / "update_knowledge_after.py").read_text(encoding="utf-8")


def set_cell_shading(cell, hex_color: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_code(doc: Document, text: str, max_chars: int = 4500):
    clipped = text if len(text) <= max_chars else text[:max_chars] + "\n# ... truncated ..."
    p = doc.add_paragraph()
    run = p.add_run(clipped)
    run.font.name = "Consolas"
    run.font.size = Pt(7)
    p.paragraph_format.space_after = Pt(8)


def build():
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("2026S-AML-2203-OTT01 — Advanced Python AI & ML Tools")
    r.bold = True
    r.font.size = Pt(12)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(
        "Final Term Project — Code Review, Quality Assessment & Refactoring\n"
        "Technical Report"
    )
    r.bold = True
    r.font.size = Pt(14)

    meta = [
        ("Course", "AML-2203 Advanced Python AI & ML Tools"),
        ("Group", "2026S-AML-2203-OTT01-Advanced Python AI & ML - 2"),
        ("Category", "Python-Groups"),
        (
            "Codebase",
            "SupportFlow / AML-2404 AI Customer Service Chatbot (Python Flask backend)",
        ),
        (
            "Primary modules reviewed",
            "backend/app.py, chatbot.py, database.py, config.py, auth.py, security.py, language_util.py",
        ),
        ("Static analysis tools", "pylint 4.0.6, flake8, black 26.5.1"),
        (
            "AI assistant used",
            "Cursor / Copilot-style AI for review suggestions and refactor design",
        ),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(meta):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
        for p in table.rows[i].cells[0].paragraphs:
            for run in p.runs:
                run.bold = True

    doc.add_heading("1. Introduction & Project Objective", level=1)
    doc.add_paragraph(
        "This report critically evaluates and improves the quality of our semester Python "
        "codebase developed for the AML-2404 / SupportFlow customer-service chatbot SaaS. "
        "The backend is a multi-module Flask application with JWT auth, SQLite persistence, "
        "Stripe billing, RAG-style knowledge retrieval, and local Ollama AI integration. "
        "The goal of AML-2203 is not to rebuild the product, but to practice clean-code "
        "habits: PEP 8 compliance, modularity, readable control flow, static analysis, "
        "and one meaningful refactor grounded in evidence from tools and manual review."
    )

    doc.add_heading("2. Selected Codebase", level=1)
    doc.add_paragraph(
        "We selected the SupportFlow Flask backend because it contains substantial logic "
        "(auth decorators, multi-tenant database helpers, billing webhooks, chat orchestration, "
        "and a large OllamaAIChatbot class). This provides enough surface area for a meaningful "
        "review of naming, duplication, complexity, and PEP issues."
    )
    for item in [
        "app.py — HTTP routes, CORS/security hooks, billing, chat endpoints",
        "chatbot.py — intent/RAG/AI response generation and sanitization",
        "database.py — SQLite schema, sessions, plan-aware public views",
        "config.py / auth.py / security.py / language_util.py — supporting modules",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3. Review Method (Manual + Tools + AI)", level=1)
    doc.add_heading("3.1 Manual review", level=2)
    doc.add_paragraph(
        "We read module boundaries and high-traffic paths (auth, knowledge update, chat "
        "prepare/stream, billing sync). We looked for duplicated business rules, unclear "
        "names, oversized functions, broad exception handlers, and places where plan limits "
        "or timestamps were redefined instead of shared."
    )
    doc.add_heading("3.2 Static analysis tools", level=2)
    doc.add_paragraph(
        "We ran flake8 (style / unused names), pylint (complexity, docs, imports), and "
        "black --check (formatter drift). Reports are stored under "
        "aml-2203-code-review/analysis/."
    )
    doc.add_heading("3.3 AI-assisted feedback", level=2)
    doc.add_paragraph(
        "Cursor AI was used to: (1) summarize hotspots in large files, (2) propose "
        "candidate refactors with lowest product risk, and (3) draft helper signatures "
        "that preserve API behavior. AI suggestions were treated as hypotheses and "
        "validated against the live route logic before changing code."
    )

    doc.add_heading("4. Summary of Tool Findings", level=1)
    findings = [
        ("Tool", "Key observations"),
        (
            "flake8",
            "17 issues across reviewed files: mostly E501 long lines; also E305 blank-line "
            "spacing in config.py, unused global in chatbot cache helper (F824), trailing "
            "blank line (W391), and one E203 slice spacing style conflict with black.",
        ),
        (
            "pylint (core modules)",
            "Rated about 9.08/10 on auth/config/security/language_util. Common notes: missing "
            "docstrings (C0116), import-outside-toplevel (C0415) for lazy imports, "
            "broad-exception-caught (W0718), too-many-return-statements in language detect.",
        ),
        (
            "pylint (app + database)",
            "Rated about 8.99/10. Hotspots: too-many-branches / statements on billing webhook "
            "and sync handlers; many broad Exception catches; logging-fstring-interpolation; "
            "missing docstrings on several Flask views.",
        ),
        (
            "black --check",
            "Would reformat language_util.py and config.py — formatting is mostly consistent "
            "but not fully black-normalized across every module.",
        ),
    ]
    ft = doc.add_table(rows=len(findings), cols=2)
    ft.style = "Table Grid"
    for i, (a, b) in enumerate(findings):
        ft.rows[i].cells[0].text = a
        ft.rows[i].cells[1].text = b
        if i == 0:
            for cell in ft.rows[i].cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(cell, "0B1F33")

    doc.add_heading("5. Written Critique — Key Observations", level=1)
    doc.add_heading("5.1 Strengths", level=2)
    doc.add_paragraph(
        "The backend already shows good modular intent: auth decorators are isolated, "
        "security helpers centralize sanitization/rate limits, config owns plan limits, "
        "and chat endpoints share _prepare_chat_request. That shared chat preparation is "
        "an example of prior successful abstraction."
    )
    doc.add_heading("5.2 Issues found (manual + tools)", level=2)
    issues = [
        "Duplicated business rules in update_knowledge — the plan word-limit error JSON and "
        "business_info defaults were copy-pasted in free-text and legacy structured branches. "
        "This risks inconsistent error messages and makes limit changes error-prone (DRY / modularity issue).",
        "Oversized route handlers — Stripe confirm/sync/webhook functions exceed pylint "
        "branch/statement budgets. Flow is correct but hard to test in isolation.",
        "PEP / style drift — long lines (E501), occasional missing blank lines (E305), and "
        "black drift in a few modules. Not runtime bugs, but reduce polish.",
        "Broad exception handling — many except Exception blocks hide root causes; acceptable "
        "at HTTP edges, weaker inside helpers.",
        "Duplicate utilities — now_utc() exists in both app.py and database.py; count_words "
        "previously imported re inside the function (fixed during cleanup).",
        "Documentation debt — pylint C0116 on many small helpers; public helpers would benefit "
        "from one-line docstrings for newcomers.",
    ]
    for i, text in enumerate(issues, start=1):
        doc.add_paragraph(f"{i}. {text}")

    doc.add_heading("5.3 Suggested improvements (prioritized)", level=2)
    for item in [
        "Extract duplicated knowledge limit / seeding helpers (implemented).",
        "Split Stripe webhook/sync into pure helpers + thin Flask views.",
        "Share a single time utility module for UTC timestamps.",
        "Narrow exception types near Stripe/API boundaries; keep generic catch only at top-level views.",
        "Run black + flake8 in CI to prevent style regressions.",
        "Add focused unit tests for plan word limits and auth decorators.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("6. Implemented Improvement — One Meaningful Refactor", level=1)
    doc.add_paragraph(
        "We chose to refactor knowledge plan-limit enforcement in backend/app.py. This is a "
        "real business rule used by paying customers’ knowledge bases. Duplication here is more "
        "dangerous than a cosmetic rename because inconsistent limits could allow over-quota "
        "saves on one code path and reject them on another."
    )
    doc.add_heading("6.1 What changed and why", level=2)
    doc.add_paragraph(
        "We introduced two helpers: _knowledge_word_limit_response(...) centralizes counting "
        "and the 403 JSON payload; _seed_knowledge_business_info(...) centralizes default "
        "business_info fields. Both free-text and legacy structured update paths now call the "
        "same helpers. Behavior stays the same for valid/invalid payloads, but future copy or "
        "limit changes happen once. We also moved import re to module scope in config.py for "
        "count_words (PEP import style)."
    )

    doc.add_heading("6.2 Before (duplicated logic)", level=2)
    add_code(doc, BEFORE)
    cap = doc.add_paragraph(
        "Figure 1. Before — identical word-limit error blocks appear twice in update_knowledge."
    )
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)

    doc.add_heading("6.3 After (shared helpers)", level=2)
    add_code(doc, AFTER)
    cap2 = doc.add_paragraph(
        "Figure 2. After — helpers own the rule; route focuses on request shape and persistence."
    )
    for run in cap2.runs:
        run.italic = True
        run.font.size = Pt(9)

    doc.add_heading("7. AI Tool Feedback Used", level=1)
    doc.add_paragraph(
        "AI review suggested several candidates: extract now_utc, split webhook handlers, "
        "and remove knowledge-limit duplication. We accepted the knowledge-limit refactor "
        "because it scored highest on: (a) clear before/after teaching value, (b) low risk "
        "to runtime behavior, and (c) direct improvement to modularity of a business rule. "
        "AI-generated helper names were edited for clarity and type of return value "
        "(word_count, optional Flask error tuple)."
    )

    doc.add_heading("8. Final Reflection — What We Learned", level=1)
    doc.add_paragraph(
        "Working on a real semester product made PEP and modularity issues more concrete "
        "than toy examples. Static tools quickly highlighted complexity and style debt, but "
        "manual review was still required to decide which findings matter for maintainability. "
        "Duplication of business rules is more costly than long lines: two identical error "
        "payloads can silently diverge. AI tools accelerate discovery and drafting, yet "
        "humans must verify that abstractions preserve API contracts. Going forward we would "
        "enforce formatter/linter checks earlier and extract helpers as soon as a rule appears "
        "a second time (Rule of Three / immediate DRY for critical limits)."
    )

    doc.add_heading("9. Folder Contents (Submission Package)", level=1)
    for item in [
        "AML-2203-Code-Review-Technical-Report.pdf — PDF version of this report",
        "AML-2203-Code-Review-Technical-Report.docx — this Word document",
        "before/update_knowledge_before.py — original excerpt",
        "after/update_knowledge_after.py — refactored excerpt",
        "analysis/ — flake8, pylint, and black outputs",
        "README.md — how to reproduce the analysis",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "The live improved code also exists in the main repository at backend/app.py and "
        "backend/config.py."
    )

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end.add_run(
        "— End of Report — Group 2026S-AML-2203-OTT01 / SupportFlow Capstone Codebase —"
    ).italic = True

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
